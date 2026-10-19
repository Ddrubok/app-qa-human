"""Export sanitized QA JSON as literal cells; requires existing openpyxl.

Usage: python export_xlsx.py report.json new-report.xlsx
Never overwrites files, installs packages, truncates cells or evaluates formulas.
"""
import argparse
import io
import json
import math
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET


def cell_value(value):
    if isinstance(value, (dict, list)) or value is None:
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, str):
        # Excel limits cell text; account for supplementary UTF-16 characters.
        if len(value.encode('utf-16-le')) // 2 > 32767:
            raise ValueError('Cell exceeds 32767 UTF-16 units; split explicitly and retain source JSON')
        if any(not (ord(c) in (9, 10, 13) or 32 <= ord(c) <= 0xD7FF
                    or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF) for c in value):
            raise ValueError('Cell contains unsupported XML characters; retain source JSON')
    elif not isinstance(value, (int, float, bool)):
        raise ValueError('Unsupported cell value')
    elif isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Non-finite numeric value')
    # Excel numeric precision: preserve long integers exactly as strings.
    elif isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
        value = str(value)
    return value


def export(report, destination):
    from openpyxl import Workbook, load_workbook
    destination = Path(destination)
    if destination.suffix.lower() != '.xlsx':
        raise ValueError('Destination must be .xlsx')
    if destination.exists():
        raise FileExistsError(destination)
    if not isinstance(report, dict):
        raise ValueError('Report must be a JSON object')
    book = Workbook()
    book.remove(book.active)
    collections = [('Cases', report.get('cases', [])), ('Coverage', report.get('coverage', []))]
    if 'findings' in report:
        collections.append(('Findings', report['findings']))
    source_rows = report.get('sources', [])
    if not isinstance(source_rows, list):
        raise ValueError('sources must be a list of records')
    context = {k: v for k, v in report.items() if k not in ('cases', 'coverage', 'sources', 'findings')}
    collections.append(('Sources', [{'kind': 'run-context', 'data': context}] + source_rows))
    expected = {}
    for name, records in collections:
        if name == 'Coverage' and isinstance(records, dict):
            records = [{'area': k, 'details': v} for k, v in records.items()]
        if not isinstance(records, list) or any(not isinstance(r, dict) for r in records):
            raise ValueError(name + ' must contain object records')
        columns = list(dict.fromkeys(key for record in records for key in record))
        if any(not isinstance(key, str) for key in columns):
            raise ValueError('Record keys must be strings')
        if len(columns) > 16384 or len(records) + 1 > 1048576:
            raise ValueError('Worksheet dimensions exceed XLSX limits')
        rows = [columns] + [[record.get(key) for key in columns] for record in records] if columns else []
        sheet = book.create_sheet(name)
        expected[name] = []
        for row_idx, row in enumerate(rows, 1):
            values = [cell_value(value) for value in row]
            expected[name].append(values)
            for col_idx, value in enumerate(values, 1):
                cell = sheet.cell(row_idx, col_idx)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                cell.hyperlink = None
        if columns:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(buffer) as archive:
        for filename in archive.namelist():
            if 'externalLinks/' in filename:
                raise ValueError('Unexpected external link')
            if filename.startswith('xl/worksheets/') and filename.endswith('.xml'):
                tree = ET.fromstring(archive.read(filename))
                if tree.find('.//' + namespace + 'f') is not None:
                    raise ValueError('Unexpected formula')
            if filename.endswith('.rels'):
                if any(node.get('TargetMode') == 'External' for node in ET.fromstring(archive.read(filename))):
                    raise ValueError('Unexpected external relationship')
    buffer.seek(0)
    reopened = load_workbook(buffer, data_only=False)
    if reopened.sheetnames != list(expected):
        raise ValueError('Sheet mismatch')
    for name, rows in expected.items():
        for i, row in enumerate(rows, 1):
            for j, value in enumerate(row, 1):
                restored = reopened[name].cell(i, j)
                # XLSX may represent an empty literal string as an empty cell.
                actual = '' if value == '' and restored.value is None else restored.value
                if actual != value or restored.data_type == 'f':
                    raise ValueError(f'Round-trip mismatch at {name}!{i},{j}')
    # All validation happens before any destination is created.
    with destination.open('xb') as output:
        output.write(buffer.getvalue())
    return {'path': str(destination), 'sheets': list(expected),
            'rows': {name: max(0, len(rows) - 1) for name, rows in expected.items()},
            'validation': 'literal values and file structure; spreadsheet UI not tested'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('destination', type=Path)
    args = parser.parse_args()
    report = json.loads(args.source.read_text(encoding='utf-8-sig'))
    print(json.dumps(export(report, args.destination), ensure_ascii=False))


if __name__ == '__main__':
    main()
